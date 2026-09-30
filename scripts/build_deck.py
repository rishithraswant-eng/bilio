#!/usr/bin/env python3
"""Build the ≤8-slide submission deck from docs/DECK_OUTLINE.md, in the app's warm theme.

    pip install python-pptx
    python3 scripts/build_deck.py [--live-score "72/100" --live-latency "0.9 s"]
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
BG, INK, MUT, ACC, RULE = (RGBColor(0x12, 0x11, 0x10), RGBColor(0xEF, 0xE8, 0xDC), RGBColor(0x9A, 0x91, 0x84),
                           RGBColor(0xE0, 0x71, 0x4A), RGBColor(0x41, 0x3B, 0x35))
W, H = Emu(12192000), Emu(6858000)


def slides(outline: str):
    for m in re.finditer(r"^\d+\. \*\*(.+?)\*\*\s*(.+)$", outline, re.M):
        yield m.group(1).rstrip("."), m.group(2).strip()


def run(p, text, size, color, font, italic=False):
    r = p.add_run()
    r.text = text
    r.font.size, r.font.color.rgb, r.font.name, r.font.italic = Pt(size), color, font, italic
    return r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--live-score", default="[pending — run ./run_fdb_v3.sh --require-judge]")
    ap.add_argument("--live-latency", default="[from results/results.md]")
    ap.add_argument("--out", default=str(ROOT / "docs/deck/BILIO_Theme05.pptx"))
    a = ap.parse_args()
    items = list(slides((ROOT / "docs/DECK_OUTLINE.md").read_text()))
    if not 1 <= len(items) <= 8:
        raise SystemExit(f"guide limit: 1-8 slides, outline has {len(items)}")
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    for i, (title, body) in enumerate(items, 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = BG
        body = body.replace("`[LIVE]` judged live pass rate and latency from `./run_fdb_v3.sh`.",
                            f"Live judged pass rate: {a.live_score}. First-response latency: {a.live_latency}.")
        body = re.sub(r"[`*]", "", body)
        k = s.shapes.add_textbox(Emu(720000), Emu(520000), Emu(9000000), Emu(400000)).text_frame
        run(k.paragraphs[0], f"{i:02d} / {len(items):02d}   BILIO · Theme 05", 11, ACC, "Courier New")
        t = s.shapes.add_textbox(Emu(720000), Emu(1000000), Emu(10700000), Emu(1500000)).text_frame
        t.word_wrap = True
        run(t.paragraphs[0], title, 40 if i == 1 else 34, INK, "Georgia", italic=(i == 1))
        line = s.shapes.add_connector(1, Emu(720000), Emu(2560000), Emu(11470000), Emu(2560000))
        line.line.color.rgb = RULE
        tf = s.shapes.add_textbox(Emu(720000), Emu(2800000), Emu(10500000), Emu(3700000)).text_frame
        tf.word_wrap = True
        parts = [x.strip() for x in re.split(r"(?<=[.;])\s+(?=[A-Z0-9(\[\"'])", body) if x.strip()]
        for j, part in enumerate(parts):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            run(p, "—  ", 17, ACC, "Helvetica")
            run(p, part, 17, INK if j == 0 else MUT, "Helvetica")
        dot = s.shapes.add_shape(9, Emu(11250000), Emu(580000), Emu(160000), Emu(160000))
        dot.fill.solid()
        dot.fill.fore_color.rgb = ACC
        dot.line.fill.background()
    Path(a.out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(a.out)
    print(f"wrote {a.out} ({len(items)} slides)")


if __name__ == "__main__":
    main()
